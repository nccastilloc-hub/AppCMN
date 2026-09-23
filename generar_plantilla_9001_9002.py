"""
generar_plantilla_9001_9002.py
Extrae las AO activas de 9001 y 9002 del POI institucional y genera
la matriz oficial para completar Definiciones Operacionales y Criterios de Programación.
"""

import os
import glob
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def buscar_archivo_poi():
    patrones = [
        "Seguimiento metas fisicas POI.xlsx",
        "Seguimiento metas fisicas POI.xls",
        "POI_*.xlsx",
        "*.xlsx"
    ]
    for p in patrones:
        archivos = glob.glob(p)
        filtrados = [a for a in archivos if "POI" in a.upper() or "METAS" in a.upper() or "SEGUIMIENTO" in a.upper()]
        if filtrados:
            filtrados.sort(key=os.path.getmtime, reverse=True)
            return filtrados[0]
        elif archivos and not p.startswith("*"):
            return archivos[0]
    return None

def main():
    archivo_origen = buscar_archivo_poi()
    if not archivo_origen:
        print("❌ No se encontró el archivo del POI en la carpeta.")
        return

    print(f"📖 Leyendo datos desde: {archivo_origen}...")
    try:
        df = pd.read_excel(archivo_origen, sheet_name="DATA")
    except Exception:
        df = pd.read_excel(archivo_origen)

    df.columns = df.columns.str.strip()

    # Filtro: AO activas pertenecientes a 9001 y 9002
    filtro = (df["Activo AO"] == "SI") & (df["Categoria ID"].isin([9001, 9002]))
    df_admin = df[filtro].copy()

    columnas_clave = [
        "Categoria ID", "Categoria",
        "Actividad Presupuestal ID", "Actividad Presupuestal",
        "Actividad Operativa ID", "Actividad Operativa",
        "Unidad de Medida"
    ]
    columnas_existentes = [c for c in columnas_clave if c in df_admin.columns]
    
    df_resumen = df_admin[columnas_existentes].drop_duplicates().sort_values(
        by=["Categoria ID", "Actividad Presupuestal ID", "Actividad Operativa ID"]
    )

    print(f"✅ Se identificaron {len(df_resumen)} actividades operativas únicas (9001 / 9002).")

    # Armado del libro Excel con formato ministerial
    wb = Workbook()
    ws = wb.active
    ws.title = "DO_CP_9001_9002"
    ws.views.sheetView[0].showGridLines = True

    # Título institucional
    ws.merge_cells("A1:K1")
    celda_titulo = ws["A1"]
    celda_titulo.value = "INSTITUTO NACIONAL MATERNO PERINATAL - DEFINICIONES OPERACIONALES Y CRITERIOS DE PROGRAMACIÓN 2026 (9001 / 9002)"
    celda_titulo.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    celda_titulo.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Cabeceras homologadas con el estándar MINSA
    headers = [
        "COD_CATEGORIA",
        "CATEGORIA PRESUPUESTAL",
        "COD_ACTIVIDAD_PRESUP",
        "ACTIVIDAD PRESUPUESTAL",
        "COD_AO_POI",
        "ACTIVIDAD OPERATIVA (DENOMINACIÓN POI)",
        "UNIDAD_MEDIDA",
        "DEFINICIÓN OPERACIONAL (DO)",
        "CRITERIO DE PROGRAMACIÓN (CP)",
        "FUENTE DE INFORMACIÓN / REGISTRO",
        "CRITERIO DE VERIFICACIÓN / CUMPLIMIENTO"
    ]

    ws.append(headers)
    ws.row_dimensions[2].height = 28

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill_fijo = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")  # Azul: Datos del POI
    header_fill_do = PatternFill(start_color="385723", end_color="385723", fill_type="solid")    # Verde: Campos a redactar

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = header_fill_do if col_idx >= 8 else header_fill_fijo

    # Poblar registros
    for _, row in df_resumen.iterrows():
        ws.append([
            row.get("Categoria ID", ""),
            row.get("Categoria", ""),
            row.get("Actividad Presupuestal ID", ""),
            row.get("Actividad Presupuestal", ""),
            row.get("Actividad Operativa ID", ""),
            row.get("Actividad Operativa", ""),
            row.get("Unidad de Medida", ""),
            "",  # Espacio para DO
            "",  # Espacio para CP
            "",  # Espacio para Fuente
            ""   # Espacio para Verificación
        ])

    # Aplicar bordes, alineación y anchos de columna
    for r in range(3, ws.max_row + 1):
        ws.row_dimensions[r].height = 24
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=9)
            if c in [1, 3, 5, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    anchos = {
        "A": 16, "B": 28, "C": 22, "D": 28,
        "E": 18, "F": 45, "G": 18, "H": 40,
        "I": 40, "J": 30, "K": 30
    }
    for col_letter, ancho in anchos.items():
        ws.column_dimensions[col_letter].width = ancho

    nombre_salida = "Matriz_DO_CP_9001_9002_INMP.xlsx"
    wb.save(nombre_salida)
    print(f"🎉 Matriz generada con éxito: {nombre_salida}")

if __name__ == "__main__":
    main()